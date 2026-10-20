"""Build the two competition documents from editable content.json.

Use the Codex bundled Python runtime (python-docx and Pillow).
All facts are source-backed in content.json and source-map.json; no credentials.
"""
from pathlib import Path
import json
from xml.sax.saxutils import escape
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
WORKSPACE = HERE.parents[2]
OUTPUT = HERE.parent / 'generated'
FONT = '/System/Library/Fonts/STHeiti Light.ttc'


class Diagram:
    def __init__(self, name, width, height):
        self.name, self.width, self.height = name, width, height
        self.im = Image.new('RGB', (width, height), 'white')
        self.d = ImageDraw.Draw(self.im)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="white"/>']
    def text(self, x, y, text, size=38, color='#172B4D'):
        font = ImageFont.truetype(FONT, size)
        for i, line in enumerate(text.split('\n')):
            yy = y + i * (size + 14)
            self.d.text((x, yy), line, font=font, anchor='mm', fill=color)
            self.svg.append(f'<text x="{x}" y="{yy+size*.34}" text-anchor="middle" font-family="Heiti SC, sans-serif" font-size="{size}" fill="{color}">{escape(line)}</text>')
    def box(self, x, y, w, h, text, fill='#F3F6FA', size=38):
        self.d.rounded_rectangle((x,y,x+w,y+h), radius=18, fill=fill, outline='#9AABC0', width=3)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="18" fill="{fill}" stroke="#9AABC0" stroke-width="3"/>')
        n = len(text.split('\n'))
        self.text(x+w/2, y+h/2-(n-1)*(size+14)/2, text, size)
    def arrow(self, x1,y1,x2,y2, label='', lx=None, ly=None):
        import math
        self.d.line((x1,y1,x2,y2), fill='#4B6381', width=4)
        angle=math.atan2(y2-y1,x2-x1)
        p=[(x2,y2),(x2-18*math.cos(angle-.45),y2-18*math.sin(angle-.45)),(x2-18*math.cos(angle+.45),y2-18*math.sin(angle+.45))]
        self.d.polygon(p,fill='#4B6381')
        self.svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#4B6381" stroke-width="4"/><polygon points="'+ ' '.join(f'{x},{y}' for x,y in p)+'" fill="#4B6381"/>')
        if label: self.text(lx if lx is not None else (x1+x2)/2,ly if ly is not None else (y1+y2)/2-30,label,31)
    def save(self):
        self.im.save(HERE / (self.name+'.png'))
        (HERE/(self.name+'.svg')).write_text('\n'.join(self.svg+['</svg>']),encoding='utf-8')


def figures():
    d=Diagram('functional-domains',2400,1220)
    for x,head,body,foot in [
        (55,'当前模拟业务','Qt SQLite 同库快照\n营收 在线率 用户数\n充电订单 完成电量\n经营趋势 异常摘要','用户操作提交后轮询可见'),
        (835,'历史运营分析','180 天模拟发布批次\n多维统计与对比\n负荷 空闲 风险与异常\n模型报告与受控 AI 解释','历史批次与来源可追溯'),
        (1615,'独立策略实验','独立 MySQL 回放环境\n预计到达与条件等待\n找站评分与激励情景\n固定请求 种子与指标','实验结果不写入 Qt 预约')]:
        d.box(x,70,730,140,head,'#EAF2FB',46)
        d.box(x,250,730,480,body,'#F3F6FA',42)
        d.box(x,770,730,130,foot,'#FFFFFF',36)
    d.box(55,1000,2290,145,'统一运营中心按来源分区展示\n当前快照、历史批次与策略情景不混算；失败保留最后成功数据并明确提示','#EAF2FB',40)
    d.save()
    d=Diagram('architecture',2400,1280)
    d.box(55,60,600,155,'Qt 用户客户端\n找站 预约 充电 支付', '#EAF2FB',42)
    d.box(895,60,610,155,'Qt 服务端与管理端\nTCP 业务服务 原生管理', '#EAF2FB',42)
    d.box(1745,60,600,155,'Qt SQLite\n当前模拟业务唯一权威', '#F3F6FA',42)
    d.arrow(655,137,895,137,'TCP',ly=100)
    d.arrow(1505,137,1745,137,'事务读写',ly=100)
    d.box(55,400,600,170,'Vue 运营中心\nQWebEngine 嵌入与全屏', '#EAF2FB',42)
    d.box(895,400,610,170,'Python FastAPI\n当前快照 历史 API 模型服务', '#EAF2FB',40)
    d.box(1745,400,600,170,'只读 SQLite 适配器\n单事务 snapshot', '#F3F6FA',42)
    d.arrow(655,485,895,485,'同源 HTTP',ly=445)
    d.arrow(1505,485,1745,485,'调用',ly=445)
    d.arrow(2045,400,2045,215,'只读',lx=2150,ly=305)
    d.arrow(1200,215,1200,400,'QProcess 自动托管',lx=1385,ly=307)
    d.box(55,790,600,180,'历史统计 MySQL\n已发布 180 天模拟批次', '#F3F6FA',42)
    d.box(895,790,610,180,'可信 CPU 模型工件\n负荷 空闲 到站 风险与异常', '#F3F6FA',40)
    d.box(1745,790,600,180,'独立策略 MySQL\n回放时钟 推荐与实验记录', '#F3F6FA',42)
    d.arrow(1200,570,355,790)
    d.arrow(1200,570,1200,790)
    d.arrow(1200,570,2045,790)
    d.box(55,1090,850,120,'腾讯地图只由客户端配置与调用', '#FFFFFF',38)
    d.box(1110,1090,1235,120,'外部 AI API 由服务端调用 仅发送受控历史上下文', '#FFFFFF',38)
    d.save()
    d=Diagram('data-flow',2400,680)
    xs=[50,650,1250,1850]
    labels=['用户 TCP 操作\n预约到模拟支付','Qt 事务提交\n更新订单与账本','只读一致快照\n汇总及明细同一时刻','Vue 轮询展示\n成功后更新时间']
    for x,t in zip(xs,labels):d.box(x,90,500,180,t,'#EAF2FB',40)
    for i in range(3):d.arrow(xs[i]+500,180,xs[i+1],180)
    d.box(60,415,1090,155,'失败分支\n保留最后成功数据 明确离线 不伪装为零', '#F3F6FA',40)
    d.box(1250,415,1090,155,'历史分支\n独立发布批次 不吸收 Qt 新增订单', '#F3F6FA',40)
    d.arrow(2100,270,2100,415,'切换历史页',lx=2260,ly=340)
    d.arrow(2100,325,605,415,'请求失败时',lx=1050,ly=340)
    d.save()


def set_font(style, size, bold=False):
    style.font.name='Heiti SC'
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.italic=False
    style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
    for attr in list(fonts.attrib):
        if attr.endswith('Theme'):del fonts.attrib[attr]
    for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),'Heiti SC')


def base_document(title):
    doc=Document();sec=doc.sections[0]
    for element in list(doc.styles.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7)
    sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    sec.footer_distance=Cm(.8)
    set_font(doc.styles['Normal'],10.5)
    p=doc.styles['Normal'].paragraph_format
    p.line_spacing=Pt(17);p.space_after=Pt(7);p.widow_control=True
    for name,size in [('Title',25),('Subtitle',14),('Heading 1',16),('Heading 2',12.5),('Caption',9)]:
        set_font(doc.styles[name],size,name!='Caption')
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].paragraph_format.space_before=Pt(11)
        doc.styles[name].paragraph_format.space_after=Pt(7)
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles['Title'].paragraph_format.space_after=Pt(10)
    doc.styles['Subtitle'].paragraph_format.space_after=Pt(12)
    set_font(doc.styles['List Bullet'],10.5)
    doc.styles['List Bullet'].paragraph_format.line_spacing=Pt(17)
    doc.styles['List Bullet'].paragraph_format.space_after=Pt(5)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('第 ');run.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    footer.add_run(' 页').font.size=Pt(9)
    doc.core_properties.title=title
    doc.core_properties.subject='2026仓颉杯综合系统赛道参赛材料'
    doc.core_properties.author='三带一对（队）'
    doc.core_properties.last_modified_by='三带一对（队）'
    doc.core_properties.keywords='充能智析 Qt Vue 运营分析 模拟数据'
    doc.core_properties.comments=''
    return doc


def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=v
    tblpr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    tblpr.append(borders)
    for ri,row in enumerate(t.rows):
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        prevent=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prevent)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            marg=OxmlElement('w:tcMar')
            for edge,val in [('top',100),('bottom',100),('left',110),('right',110)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');marg.append(el)
            tcpr.append(marg)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'284968' if ri==0 else ('F3F6F9' if ri%2==0 else 'FFFFFF'));tcpr.append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(14.5)
                if ci==0 and len(headers)>2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9.5);r.font.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def build(item):
    doc=base_document(item['title'])
    for pi,page in enumerate(item['pages']):
        if pi:doc.add_page_break()
        for block in page:
            kind=block['type']
            if kind in ['title','subtitle','h1','h2','p','bullet','caption']:
                style={'title':'Title','subtitle':'Subtitle','h1':'Heading 1','h2':'Heading 2','p':'Normal','bullet':'List Bullet','caption':'Caption'}[kind]
                p=doc.add_paragraph(block['text'],style)
                if kind=='caption':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            elif kind=='table':table(doc,block['headers'],block['rows'],block['widths'])
            elif kind=='link':
                p=doc.add_paragraph()
                p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.line_spacing=Pt(14)
                link=OxmlElement('w:hyperlink')
                link.set(qn('r:id'),doc.part.relate_to(block['url'],RT.HYPERLINK,is_external=True))
                run=OxmlElement('w:r');props=OxmlElement('w:rPr')
                size=OxmlElement('w:sz');size.set(qn('w:val'),'18');props.append(size)
                color=OxmlElement('w:color');color.set(qn('w:val'),'284968');props.append(color)
                run.append(props);text=OxmlElement('w:t');text.text=block['text'];run.append(text)
                link.append(run);p._p.append(link)
            elif kind=='image':
                path=HERE/block['path'] if block.get('local') else WORKSPACE/block['path']
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True
                p.paragraph_format.line_spacing=1.0
                p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.space_after=Pt(4)
                p.add_run().add_picture(str(path),width=Cm(block.get('width',17)))
                inline=p.runs[0]._r.xpath('.//wp:docPr')[0];inline.set('descr',block['alt'])
    OUTPUT.mkdir(parents=True,exist_ok=True)
    path=OUTPUT/item['file'];doc.save(path)
    return path


if __name__=='__main__':
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=OUTPUT,
                        help='Directory for the two generated DOCX files')
    parser.add_argument('--content', type=Path, default=HERE/'content.json',
                        help='Editable content JSON file')
    args = parser.parse_args()
    OUTPUT = args.output_dir.resolve()
    figures()
    contents=json.loads(args.content.read_text(encoding='utf-8'))
    for item in contents: print(build(item))
